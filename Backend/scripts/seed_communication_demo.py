"""
Create a dummy careers application and optional parser-style data for communication tests.

What it does
------------
1. Builds a small **.docx** resume (same formats as ``POST /api/careers/apply``).
2. Calls ``apply_for_job_service`` (Careers module) — same logic as the public apply API.
3. Then either:
   - **``--synthetic-shortlist``** (default): sets ``screening_status=shortlisted``, scores, and
     ``email`` so you can call communication APIs without OpenAI for screening.
   - **``--run-openai-parser``**: runs ``run_screening_for_candidate_id`` (Parser module) — needs
     ``OPENAI_API_KEY`` and enough resume text vs JD.

Usage
-----
  cd Backend
  python -m scripts.seed_communication_demo --jd-id <uuid>

  # Resolve JD from an onboarded user's email (newest JD in their org)
  python -m scripts.seed_communication_demo --owner-email you@company.com

  # Real parser (OpenAI) instead of synthetic fields
  python -m scripts.seed_communication_demo --jd-id <uuid> --run-openai-parser --no-synthetic-shortlist

  # Apply only (status stays ``pending``)
  python -m scripts.seed_communication_demo --jd-id <uuid> --no-synthetic-shortlist

Then test communication (JWT = that org's owner):

  POST /api/communication/candidates/{candidate_id}/emails/shortlist
  POST /api/communication/candidates/{candidate_id}/voice-interviews/schedule
"""

from __future__ import annotations

import argparse
import json
import sys
from datetime import datetime
from io import BytesIO

from sqlalchemy.orm import Session

from database.database import SessionLocal
from database.models import Candidate, JobDescription, Users
from modules.Careers.careers_service import apply_for_job_service
from modules.parser.parser_service import run_screening_for_candidate_id
from utils.organization_utils import get_organization_for_owner


def _demo_resume_docx_bytes() -> bytes:
    """Text-heavy .docx so optional OpenAI parser has material (and extract_text succeeds)."""
    from docx import Document

    doc = Document()
    doc.add_heading("Alex Rivera", 0)
    doc.add_paragraph(
        "Email: demo.communication@test.local | Phone: +91-9876543210 | Location: Bengaluru"
    )
    doc.add_heading("Professional Summary", level=1)
    doc.add_paragraph(
        "Software engineer with 4+ years building responsive web applications using "
        "React, TypeScript, and modern JavaScript. Strong experience with REST APIs, "
        "Git workflows, and agile delivery. Passionate about performance, accessibility, "
        "and clean UI engineering."
    )
    doc.add_heading("Skills", level=1)
    doc.add_paragraph(
        "JavaScript, TypeScript, React, HTML5, CSS3, Redux, Node.js basics, Git, Jest, "
        "Webpack, REST APIs, responsive design, cross-browser compatibility."
    )
    doc.add_heading("Experience", level=1)
    doc.add_paragraph(
        "Senior Frontend Developer — TechCorp (2021–Present)\n"
        "• Built customer-facing dashboards with React and reusable component libraries.\n"
        "• Collaborated with product and design to ship A/B tested UX improvements.\n"
        "• Improved Core Web Vitals through code-splitting and lazy loading.\n"
    )
    doc.add_paragraph(
        "Frontend Developer — StartupLabs (2019–2021)\n"
        "• Developed marketing sites and internal tools with Vue.js and React.\n"
        "• Integrated third-party APIs and implemented role-based access in the UI.\n"
    )
    doc.add_heading("Education", level=1)
    doc.add_paragraph("B.Tech Computer Science — State University (2019)")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _resolve_jd_id(db: Session, jd_id: str, owner_email: str) -> tuple[str | None, str | None]:
    jid = (jd_id or "").strip()
    if jid:
        row = db.query(JobDescription).filter(JobDescription.id == jid).first()
        if not row:
            return None, f"Job description not found: {jid}"
        return jid, None
    email = (owner_email or "").strip().lower()
    if not email:
        return None, "Pass --jd-id or --owner-email"
    user = db.query(Users).filter(Users.email == email).first()
    if not user:
        return None, f"No user with email: {email}"
    org = get_organization_for_owner(db, user.id)
    if not org:
        return None, "User has no organization (complete onboarding first)."
    jd = (
        db.query(JobDescription)
        .filter(JobDescription.organization_id == org.id)
        .order_by(JobDescription.created_at.desc())
        .first()
    )
    if not jd:
        return None, "No job descriptions for this organization. Use /api/broadcasting/generate-jd first."
    return jd.id, None


def _synthetic_shortlist(db: Session, candidate_id: str, email: str) -> None:
    row = db.query(Candidate).filter(Candidate.id == candidate_id).first()
    if not row:
        raise ValueError("Candidate missing after apply")
    now = datetime.utcnow()
    details = {
        "strengths_vs_jd": "Strong React/JS match; demo seed data for communication tests.",
        "gaps_vs_jd": "None flagged in synthetic data.",
        "seed_script": "scripts.seed_communication_demo",
    }
    row.email = (email or "").strip().lower() or row.email
    row.screening_status = "shortlisted"
    row.ats_match_score = 85.0
    row.jd_match_analysis = (
        "**Synthetic screening (no OpenAI run).** Candidate marked shortlisted for "
        "testing shortlist email and Retell voice flows."
    )
    row.screening_details_json = json.dumps(details, ensure_ascii=False)
    row.parsed_resume_summary = (
        "Demo candidate: frontend engineer, React/TypeScript, 4+ years (seed data)."
    )
    row.parser_error_message = None
    row.screening_updated_at = now
    row.updated_at = now
    db.commit()


def main() -> int:
    p = argparse.ArgumentParser(description="Seed careers + parser data for communication tests")
    p.add_argument("--jd-id", default="", help="Target job description UUID")
    p.add_argument(
        "--owner-email",
        default="",
        help="Onboarded user email; uses their org's newest JD if --jd-id omitted",
    )
    p.add_argument("--name", default="Alex Rivera (Communication Demo)")
    p.add_argument("--phone", default="9876543210")
    p.add_argument(
        "--email",
        default="demo.communication@test.local",
        help="Stored on candidate for shortlist / invite emails",
    )
    p.add_argument(
        "--run-openai-parser",
        action="store_true",
        help="Run real parser (OpenAI) after apply",
    )
    p.add_argument(
        "--synthetic-shortlist",
        action=argparse.BooleanOptionalAction,
        default=True,
        help="Set shortlisted + dummy scores without OpenAI (default: true)",
    )
    args = p.parse_args()

    db = SessionLocal()
    try:
        jd_id, err = _resolve_jd_id(db, args.jd_id, args.owner_email)
        if err:
            print(err, file=sys.stderr)
            return 1
        assert jd_id is not None
    finally:
        db.close()

    docx = _demo_resume_docx_bytes()
    apply_out = apply_for_job_service(
        jd_id,
        args.name,
        args.phone,
        docx,
        "demo_communication_seed.docx",
    )
    print(json.dumps(apply_out, indent=2))
    if not apply_out.get("status"):
        return 1

    candidate_id = (apply_out.get("data") or {}).get("id")
    if not candidate_id:
        print("Apply succeeded but no candidate id in response", file=sys.stderr)
        return 1

    if args.run_openai_parser:
        print("\nRunning OpenAI parser (run_screening_for_candidate_id)...", file=sys.stderr)
        run_screening_for_candidate_id(candidate_id)
        db2 = SessionLocal()
        try:
            row = db2.query(Candidate).filter(Candidate.id == candidate_id).first()
            if row and (args.email or "").strip():
                row.email = args.email.strip().lower()
                row.updated_at = datetime.utcnow()
                db2.commit()
        finally:
            db2.close()
        fresh = SessionLocal()
        try:
            row = fresh.query(Candidate).filter(Candidate.id == candidate_id).first()
            print(
                "\nAfter parser:",
                json.dumps(
                    {
                        "id": row.id,
                        "screening_status": row.screening_status,
                        "ats_match_score": row.ats_match_score,
                        "email": row.email,
                    },
                    indent=2,
                ),
            )
        finally:
            fresh.close()

    elif args.synthetic_shortlist:
        db3 = SessionLocal()
        try:
            _synthetic_shortlist(db3, candidate_id, args.email)
        finally:
            db3.close()
        print(
            "\nSynthetic shortlist applied.",
            f"\n  candidate_id={candidate_id}",
            "\n  With employer JWT, try:",
            f"\n    POST /api/communication/candidates/{candidate_id}/emails/shortlist",
            f"\n    POST /api/communication/candidates/{candidate_id}/voice-interviews/schedule",
            file=sys.stderr,
        )
    else:
        print(
            f"\nCandidate {candidate_id} left as pending (no parser). "
            "POST /api/parser/candidates/{id}/analyze with JWT or re-run with --synthetic-shortlist.",
            file=sys.stderr,
        )

    print(f"\ncandidate_id={candidate_id}", file=sys.stderr)
    print(f"jd_id={jd_id}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
